import io
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session
from typing import List, Tuple
from app.models import Student
from app.schemas import ImportResult, ImportRowError
from app import crud

# Ánh xạ các tên cột phổ biến sang trường dữ liệu chuẩn
COLUMN_MAPPING = {
    "mã sinh viên": "student_code",
    "ma sinh vien": "student_code",
    "mã sv": "student_code",
    "ma sv": "student_code",
    "student_code": "student_code",
    "student code": "student_code",
    "mssv": "student_code",

    "họ và tên": "full_name",
    "ho va ten": "full_name",
    "họ tên": "full_name",
    "ho ten": "full_name",
    "tên": "full_name",
    "full_name": "full_name",
    "fullname": "full_name",

    "email": "email",
    "thư điện tử": "email",

    "số điện thoại": "phone",
    "so dien thoai": "phone",
    "sđt": "phone",
    "sdt": "phone",
    "phone": "phone",

    "ngày sinh": "date_of_birth",
    "ngay sinh": "date_of_birth",
    "dob": "date_of_birth",
    "date_of_birth": "date_of_birth",

    "giới tính": "gender",
    "gioi tinh": "gender",
    "gender": "gender",

    "chuyên ngành": "major",
    "chuyen nganh": "major",
    "ngành": "major",
    "nganh": "major",
    "khoa": "major",
    "major": "major",

    "điểm trung bình": "gpa",
    "diem trung binh": "gpa",
    "điểm tb": "gpa",
    "diem tb": "gpa",
    "đtb": "gpa",
    "dtb": "gpa",
    "gpa": "gpa"
}

def generate_template_excel() -> io.BytesIO:
    """Tạo file Excel mẫu để người dùng tải về nhập liệu"""
    wb = Workbook()
    ws = wb.active
    ws.title = "Mau_Nhap_Sinh_Vien"

    headers = [
        "Mã sinh viên (*)",
        "Họ và tên (*)",
        "Email",
        "Số điện thoại",
        "Ngày sinh (YYYY-MM-DD)",
        "Giới tính (Nam/Nữ)",
        "Chuyên ngành (*)",
        "Điểm trung bình (0-10)"
    ]
    ws.append(headers)

    # Style Header
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    thin_border = Border(
        left=Side(style='thin', color='D3D3D3'),
        right=Side(style='thin', color='D3D3D3'),
        top=Side(style='thin', color='D3D3D3'),
        bottom=Side(style='thin', color='D3D3D3')
    )

    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = thin_border

    # Mẫu dữ liệu tham khảo
    sample_rows = [
        ["SV001", "Nguyễn Văn An", "an.nguyen@example.com", "0912345678", "2003-05-15", "Nam", "Công nghệ thông tin", 8.5],
        ["SV002", "Trần Thị Bích", "bich.tran@example.com", "0987654321", "2003-08-20", "Nữ", "Khoa học máy tính", 9.0],
        ["SV003", "Lê Hoàng Long", "long.le@example.com", "0905123456", "2003-11-02", "Nam", "Hệ thống thông tin", 7.8],
        ["SV004", "Phạm Minh Thư", "thu.pham@example.com", "0934567890", "2003-02-28", "Nữ", "Thương mại điện tử", 8.2],
    ]

    for row in sample_rows:
        ws.append(row)

    # Căn chỉnh kích thước cột
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 5, 16)

    # Đặt độ cao hàng
    ws.row_dimensions[1].height = 28

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

def parse_import_file(file_bytes: bytes, filename: str, db: Session) -> ImportResult:
    """Đọc dữ liệu từ file Excel hoặc CSV và lưu vào cơ sở dữ liệu"""
    errors: List[ImportRowError] = []
    success_count = 0

    try:
        if filename.endswith(".csv"):
            try:
                # Đọc CSV hỗ trợ utf-8-sig hoặc utf-8
                df = pd.read_csv(io.BytesIO(file_bytes), encoding="utf-8-sig")
            except UnicodeDecodeError:
                df = pd.read_csv(io.BytesIO(file_bytes), encoding="latin1")
        else:
            df = pd.read_excel(io.BytesIO(file_bytes))
    except Exception as e:
        return ImportResult(
            total_rows=0,
            success_count=0,
            error_count=1,
            errors=[ImportRowError(row=0, error=f"Không thể đọc file: {str(e)}")]
        )

    # Chuẩn hóa tên cột
    cleaned_columns = {}
    for col in df.columns:
        col_clean = str(col).strip().lower().replace("(*)", "").strip()
        if col_clean in COLUMN_MAPPING:
            cleaned_columns[col] = COLUMN_MAPPING[col_clean]

    df.rename(columns=cleaned_columns, inplace=True)

    required_fields = ["student_code", "full_name", "major"]
    missing_required = [f for f in required_fields if f not in df.columns]
    if missing_required:
        return ImportResult(
            total_rows=len(df),
            success_count=0,
            error_count=len(df),
            errors=[ImportRowError(row=1, error=f"File thiếu các cột bắt buộc: {', '.join(missing_required)}")]
        )

    existing_codes_in_db = {s.student_code for s in db.query(Student.student_code).all()}
    seen_in_file = set()

    for idx, row in df.iterrows():
        row_num = idx + 2  # Tính theo dòng hiển thị trên Excel (Header là dòng 1)
        raw_code = str(row.get("student_code", "")).strip()

        if pd.isna(row.get("student_code")) or not raw_code or raw_code.lower() == "nan":
            errors.append(ImportRowError(row=row_num, error="Mã sinh viên không được để trống."))
            continue

        if raw_code in seen_in_file:
            errors.append(ImportRowError(row=row_num, student_code=raw_code, error="Mã sinh viên bị trùng trong file nhập."))
            continue
        seen_in_file.add(raw_code)

        if raw_code in existing_codes_in_db:
            errors.append(ImportRowError(row=row_num, student_code=raw_code, error="Mã sinh viên đã tồn tại trong cơ sở dữ liệu."))
            continue

        raw_name = str(row.get("full_name", "")).strip()
        if pd.isna(row.get("full_name")) or not raw_name or raw_name.lower() == "nan":
            errors.append(ImportRowError(row=row_num, student_code=raw_code, error="Họ và tên không được để trống."))
            continue

        raw_major = str(row.get("major", "")).strip()
        if pd.isna(row.get("major")) or not raw_major or raw_major.lower() == "nan":
            errors.append(ImportRowError(row=row_num, student_code=raw_code, error="Chuyên ngành không được để trống."))
            continue

        # Điểm GPA
        gpa_val = 0.0
        if "gpa" in df.columns and not pd.isna(row.get("gpa")):
            try:
                gpa_val = float(row.get("gpa"))
                if gpa_val < 0.0 or gpa_val > 10.0:
                    errors.append(ImportRowError(row=row_num, student_code=raw_code, error="Điểm trung bình phải từ 0.0 đến 10.0."))
                    continue
            except (ValueError, TypeError):
                errors.append(ImportRowError(row=row_num, student_code=raw_code, error="Điểm trung bình không hợp lệ."))
                continue

        # Ngày sinh
        dob_val = None
        if "date_of_birth" in df.columns and not pd.isna(row.get("date_of_birth")):
            dob_val = str(row.get("date_of_birth")).strip()
            # Xử lý nếu pandas đọc thành datetime timestamp
            if " " in dob_val:
                dob_val = dob_val.split(" ")[0]

        # Email & Phone & Giới tính
        email_val = None
        if "email" in df.columns and not pd.isna(row.get("email")):
            raw_email = str(row.get("email")).strip()
            if raw_email and raw_email.lower() != "nan":
                email_val = raw_email

        phone_val = None
        if "phone" in df.columns and not pd.isna(row.get("phone")):
            raw_phone = str(row.get("phone")).strip()
            # Tránh trường hợp đọc số điện thoại dạng float .0
            if raw_phone.endswith(".0"):
                raw_phone = raw_phone[:-2]
            if raw_phone and raw_phone.lower() != "nan":
                phone_val = raw_phone

        gender_val = "Nam"
        if "gender" in df.columns and not pd.isna(row.get("gender")):
            raw_gender = str(row.get("gender")).strip()
            if raw_gender and raw_gender.lower() != "nan":
                gender_val = raw_gender

        # Tạo record mới
        new_student = Student(
            student_code=raw_code,
            full_name=raw_name,
            email=email_val,
            phone=phone_val,
            date_of_birth=dob_val,
            gender=gender_val,
            major=raw_major,
            gpa=gpa_val
        )
        db.add(new_student)
        existing_codes_in_db.add(raw_code)
        success_count += 1

    db.commit()

    return ImportResult(
        total_rows=len(df),
        success_count=success_count,
        error_count=len(errors),
        errors=errors
    )

def export_students_to_excel(students: List[Student]) -> io.BytesIO:
    """Xuất danh sách sinh viên ra file Excel định dạng đẹp"""
    wb = Workbook()
    ws = wb.active
    ws.title = "Danh_Sach_Sinh_Vien"

    # Tiêu đề bảng lớn
    ws.merge_cells("A1:H1")
    title_cell = ws["A1"]
    title_cell.value = "DANH SÁCH SINH VIÊN"
    title_cell.font = Font(name="Arial", size=16, bold=True, color="1F4E79")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 35

    # Dòng Header
    headers = ["Mã SV", "Họ và tên", "Email", "Số điện thoại", "Ngày sinh", "Giới tính", "Chuyên ngành", "Điểm TB"]
    ws.append([]) # Dòng 2 để trống tạo khoảng cách
    ws.append(headers) # Dòng 3

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='BFBFBF'),
        right=Side(style='thin', color='BFBFBF'),
        top=Side(style='thin', color='BFBFBF'),
        bottom=Side(style='thin', color='BFBFBF')
    )

    for col_idx in range(1, len(headers) + 1):
        c = ws.cell(row=3, column=col_idx)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border

    ws.row_dimensions[3].height = 26

    # Điền dữ liệu sinh viên
    for student in students:
        ws.append([
            student.student_code,
            student.full_name,
            student.email or "",
            student.phone or "",
            student.date_of_birth or "",
            student.gender or "",
            student.major,
            student.gpa
        ])

    # Định dạng dữ liệu và kẻ khung
    for row in ws.iter_rows(min_row=4, max_row=ws.max_row, min_col=1, max_col=8):
        for cell in row:
            cell.border = thin_border
            cell.font = Font(name="Arial", size=10)
            if cell.column in [1, 5, 6, 8]:  # Mã SV, Ngày sinh, Giới tính, Điểm
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    # Tự động căn chỉnh độ rộng cột
    for col in ws.columns:
        max_len = 0
        for cell in col:
            if cell.row > 1 and cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 5, 14)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

def export_students_to_csv(students: List[Student]) -> bytes:
    """Xuất danh sách sinh viên ra định dạng CSV chuẩn UTF-8 có BOM"""
    data = []
    for s in students:
        data.append({
            "Mã sinh viên": s.student_code,
            "Họ và tên": s.full_name,
            "Email": s.email or "",
            "Số điện thoại": s.phone or "",
            "Ngày sinh": s.date_of_birth or "",
            "Giới tính": s.gender or "",
            "Chuyên ngành": s.major,
            "Điểm trung bình": s.gpa
        })

    df = pd.DataFrame(data)
    csv_str = df.to_csv(index=False, encoding="utf-8-sig")
    return csv_str.encode("utf-8-sig")
